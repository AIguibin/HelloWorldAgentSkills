# -*- coding: utf-8 -*-
"""
Task 5: 生成最终分区清单 Excel
合并 AI 分析结果 + 人工标注基线（责任人信息）
"""
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = r"e:\WorkSpace\HelloWorldAgentSkills\aiguibin-common-excel\DD-湖北农信基础配置检核"

# ==================== 加载数据 ====================

# 1. AI分析结果
with open(os.path.join(BASE_DIR, "ai_analysis_result.json"), "r", encoding="utf-8") as f:
    ai_result = json.load(f)

# 2. 原始分区基线（获取责任人、所属数据库信息）
with open(os.path.join(BASE_DIR, "partition_analysis_data.json"), "r", encoding="utf-8") as f:
    baseline_data = json.load(f)

# 3. 表结构（获取所属数据库信息）
with open(os.path.join(BASE_DIR, "table_structures.json"), "r", encoding="utf-8") as f:
    structures_data = json.load(f)

# ==================== 构建基线映射 ====================

# 从原始分区基线构建 {表英文名(小写): {责任人, 所属数据库}} 映射
baseline_map = {}
for item in baseline_data.get("existing_partitions", []):
    en = item.get("表名（英文）", "").lower()
    baseline_map[en] = {
        "责任人": item.get("责任人", "待确认"),
        "所属数据库": item.get("所属数据库", ""),
    }

# 模块 → 数据库的默认映射（从原始基线推断）
MODULE_DB_MAP = {
    "合同管理": "NCMS_CREDIT_BASE",
    "客户管理": "NCMS_CUSTOMER_BASE",
    "客户中心库": "NCMS_CUSTOMER_BASE",
    "押品管理": "NCMS_CREDIT_BASE",
    "贷后管理": "NCMS_CREDIT_BASE",
    "用信管理": "NCMS_CREDIT_BASE",
    "授信管理": "NCMS_CREDIT_BASE",
    "额度中心库": "NCMS_CREDIT_BASE",
    "产品管理": "NCMS_CREDIT_BASE",
    "档案管理": "NCMS_CREDIT_BASE",
    "放还款组": "NCMS_CREDIT_BASE",
    "线上贷款": "NCMS_CREDIT_BASE",
    "营销管理": "NCMS_CREDIT_BASE",
    "评级管理": "NCMS_CREDIT_BASE",
    "系统管理": "NCMS_CREDIT_BASE",
    "工作流程库": "NCMS_CREDIT_BASE",
    "批量管理": "NCMS_CREDIT_BASE",
    "架构管理": "NCMS_CREDIT_BASE",
    "风控中心": "NCMS_CREDIT_BASE",
    "风控中心库": "NCMS_CREDIT_BASE",
    "信贷引擎库": "NCMS_CREDIT_BASE",
    "电子文档库": "NCMS_CREDIT_BASE",
    "统一认证库": "NCMS_CREDIT_BASE",
}

# 从原始基线统计每个模块的数据库
module_db_count = {}
for item in baseline_data.get("existing_partitions", []):
    module = item.get("所属模块", "")
    db = item.get("所属数据库", "")
    if module and db:
        if module not in module_db_count:
            module_db_count[module] = {}
        module_db_count[module][db] = module_db_count[module].get(db, 0) + 1

# 更新模块默认数据库（取出现最多的）
for module, db_counts in module_db_count.items():
    most_common_db = max(db_counts, key=db_counts.get)
    MODULE_DB_MAP[module] = most_common_db

# ==================== 生成 Excel ====================

wb = Workbook()
ws = wb.active
ws.title = "分区清单"

# 样式定义
title_font = Font(name="微软雅黑", size=14, bold=True)
header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
data_font = Font(name="微软雅黑", size=9)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
partition_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # 浅黄色标记分区表
thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

# 列宽
col_widths = {
    "A": 6,   # 序号
    "B": 20,  # 所属数据库
    "C": 14,  # 所属模块
    "D": 10,  # 责任人
    "E": 30,  # 表名（英文）
    "F": 30,  # 表名（中文）
    "G": 12,  # 是否表分区
    "H": 20,  # 表分区键
    "I": 14,  # 初始表分区数
    "J": 45,  # 表分区策略
    "K": 60,  # 备注
}
for col, width in col_widths.items():
    ws.column_dimensions[col].width = width

# 标题行（第1行）
ws.merge_cells("A1:K1")
title_cell = ws["A1"]
title_cell.value = "湖北农信新信贷数据库表分区清单（V3 AI语义分析版）"
title_cell.font = title_font
title_cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 30

# 表头行（第2行）
headers = ["序号", "所属数据库", "所属模块", "责任人", "表名（英文）", "表名（中文）", "是否表分区", "表分区键", "初始表分区数", "表分区策略", "备注"]
for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=2, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align
    cell.border = thin_border
ws.row_dimensions[2].height = 25

# 数据行
# 排序：先分区表（按预估数据量降序），再不分区表（按模块+表名排序）
ai_tables = ai_result["tables"]
partition_tables = sorted([t for t in ai_tables if t["need_partition"]], key=lambda x: -x["estimated_volume"])
no_partition_tables = sorted([t for t in ai_tables if not t["need_partition"]], key=lambda x: (x["module"], x["table_en"]))

all_tables = partition_tables + no_partition_tables

row = 3
for idx, table in enumerate(all_tables):
    seq = idx + 1
    table_en = table["table_en"]
    table_en_lower = table_en.lower()
    module = table["module"]
    
    # 获取责任人
    if table_en_lower in baseline_map:
        person = baseline_map[table_en_lower]["责任人"]
        db = baseline_map[table_en_lower]["所属数据库"]
    else:
        person = "待确认"
        db = MODULE_DB_MAP.get(module, "NCMS_CREDIT_BASE")
    
    # 是否分区
    if table["need_partition"]:
        is_partition = "是"
        partition_key = table["partition_key"]
        partition_count = str(table["partition_count"])
        strategy = table["partition_strategy"]
        remark = table["remark"]
        fill = partition_fill  # 浅黄色标记
    else:
        is_partition = "否"
        partition_key = ""
        partition_count = ""
        strategy = ""
        remark = table["remark"]
        fill = None
    
    values = [
        seq, db, module, person, table_en, table["table_cn"],
        is_partition, partition_key, partition_count, strategy, remark
    ]
    
    for col_idx, value in enumerate(values, 1):
        cell = ws.cell(row=row, column=col_idx, value=value)
        cell.font = data_font
        cell.border = thin_border
        if col_idx in [1, 2, 3, 4, 7, 8, 9]:  # 居中对齐
            cell.alignment = center_align
        else:  # 左对齐
            cell.alignment = left_align
        if fill:
            cell.fill = fill
    
    ws.row_dimensions[row].height = 22
    row += 1

# 冻结窗格
ws.freeze_panes = "A3"

# 添加筛选
ws.auto_filter.ref = f"A2:K{row - 1}"

# 保存
output_path = os.path.join(BASE_DIR, "湖北农信数据库分区表清单_精细化版_v3.xlsx")
wb.save(output_path)
print(f"Excel已生成: {output_path}")
print(f"总行数: {len(all_tables)}")
print(f"分区表: {len(partition_tables)}")
print(f"不分区表: {len(no_partition_tables)}")