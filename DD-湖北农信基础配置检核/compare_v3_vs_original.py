# -*- coding: utf-8 -*-
"""
对比 V3 版分区清单 vs 原始分区清单
"""
import json
import os
from openpyxl import load_workbook

BASE_DIR = r"e:\WorkSpace\HelloWorldAgentSkills\aiguibin-common-excel\DD-湖北农信基础配置检核"

# ==================== 读取原始清单 ====================
print("=" * 70)
print("读取原始分区清单...")
wb_orig = load_workbook(os.path.join(BASE_DIR, "湖北农信数据库分区表清单.xlsx"))
ws_orig = wb_orig["数据库表分区"]

# 第2行是表头，数据从第3行开始
orig_tables = {}  # {表英文名(小写): {完整信息}}
for row in ws_orig.iter_rows(min_row=3, values_only=True):
    if row[4] is None:
        continue
    table_en = str(row[4]).strip().lower()
    is_partition = str(row[6]).strip() if row[6] else ""
    orig_tables[table_en] = {
        "序号": row[0],
        "所属数据库": row[1],
        "所属模块": row[2],
        "责任人": row[3],
        "表名（英文）": row[4],
        "表名（中文）": row[5],
        "是否表分区": is_partition,
        "表分区键": row[7] if row[7] else "",
        "初始表分区数": row[8] if row[8] else "",
        "表分区策略": row[9] if row[9] else "",
        "备注": row[10] if row[10] else "",
    }

orig_partition = {k: v for k, v in orig_tables.items() if v["是否表分区"] == "是"}
print(f"原始清单: 总{len(orig_tables)}张表, 分区{len(orig_partition)}张")

# ==================== 读取 V3 清单 ====================
print("读取 V3 分区清单...")
wb_v3 = load_workbook(os.path.join(BASE_DIR, "湖北农信数据库分区表清单_精细化版_v3.xlsx"))
ws_v3 = wb_v3["分区清单"]

v3_tables = {}
for row in ws_v3.iter_rows(min_row=3, values_only=True):
    if row[4] is None:
        continue
    table_en = str(row[4]).strip().lower()
    v3_tables[table_en] = {
        "序号": row[0],
        "所属数据库": row[1],
        "所属模块": row[2],
        "责任人": row[3],
        "表名（英文）": row[4],
        "表名（中文）": row[5],
        "是否表分区": str(row[6]).strip() if row[6] else "",
        "表分区键": row[7] if row[7] else "",
        "初始表分区数": row[8] if row[8] else "",
        "表分区策略": row[9] if row[9] else "",
        "备注": row[10] if row[10] else "",
    }

v3_partition = {k: v for k, v in v3_tables.items() if v["是否表分区"] == "是"}
print(f"V3 清单: 总{len(v3_tables)}张表, 分区{len(v3_partition)}张")

# ==================== 差异分析 ====================

# 原始标记为分区、V3标记为不分区 → 被移除
removed = {}
for en, info in orig_partition.items():
    if en in v3_tables:
        if v3_tables[en]["是否表分区"] != "是":
            removed[en] = {
                "表中文名": info["表名（中文）"],
                "表英文名": info["表名（英文）"],
                "模块": info["所属模块"],
                "原始分区键": info["表分区键"],
                "原始分区数": info["初始表分区数"],
                "V3判断": "不分区",
            }
    else:
        # 原始清单有，V3清单没有这张表
        removed[en] = {
            "表中文名": info["表名（中文）"],
            "表英文名": info["表名（英文）"],
            "模块": info["所属模块"],
            "原始分区键": info["表分区键"],
            "原始分区数": info["初始表分区数"],
            "V3判断": "V3清单中不存在此表",
        }

# V3标记为分区、原始标记为不分区（或不存在）→ 新加入
newly_added = {}
for en, info in v3_partition.items():
    if en in orig_tables:
        if orig_tables[en]["是否表分区"] != "是":
            newly_added[en] = {
                "表中文名": info["表名（中文）"],
                "表英文名": info["表名（英文）"],
                "模块": info["所属模块"],
                "V3分区键": info["表分区键"],
                "V3分区数": info["初始表分区数"],
                "原始判断": "不分区",
            }
    else:
        # V3新增的表（原始清单中不存在）
        newly_added[en] = {
            "表中文名": info["表名（中文）"],
            "表英文名": info["表名（英文）"],
            "模块": info["所属模块"],
            "V3分区键": info["表分区键"],
            "V3分区数": info["初始表分区数"],
            "原始判断": "原始清单中不存在",
        }

# 两者都标记为分区 → 对比分区键/分区数变化
both_partition = {}
for en, info in v3_partition.items():
    if en in orig_partition:
        orig_key = str(orig_partition[en]["表分区键"]).strip()
        v3_key = str(info["表分区键"]).strip()
        orig_count = str(orig_partition[en]["初始表分区数"]).strip()
        v3_count = str(info["初始表分区数"]).strip()
        
        if orig_key != v3_key or orig_count != v3_count:
            both_partition[en] = {
                "表中文名": info["表名（中文）"],
                "表英文名": info["表名（英文）"],
                "模块": info["所属模块"],
                "原始分区键": orig_key,
                "V3分区键": v3_key,
                "原始分区数": orig_count,
                "V3分区数": v3_count,
            }

# ==================== 输出对比结果 ====================
print("\n" + "=" * 70)
print("对比结果汇总")
print("=" * 70)
print(f"原始清单分区表: {len(orig_partition)} 张")
print(f"V3 清单分区表:   {len(v3_partition)} 张")
print(f"被移除:          {len(removed)} 张")
print(f"新加入:          {len(newly_added)} 张")
print(f"分区键/分区数变化: {len(both_partition)} 张")
print(f"保持一致:        {len(v3_partition) - len(newly_added) - len(both_partition)} 张")

# ==================== 1. 被移除的表 ====================
print("\n" + "=" * 70)
print(f"一、被移除的分区表（{len(removed)} 张）—— 原始标记为分区，V3标记为不分区")
print("=" * 70)

if removed:
    # 按模块分组
    by_module = {}
    for en, info in removed.items():
        mod = info["模块"]
        if mod not in by_module:
            by_module[mod] = []
        by_module[mod].append(info)
    
    for mod in sorted(by_module.keys()):
        tables = by_module[mod]
        print(f"\n  [{mod}] {len(tables)}张:")
        for t in tables:
            orig_key = t.get("原始分区键", "")
            orig_cnt = t.get("原始分区数", "")
            print(f"    - {t['表中文名']}({t['表英文名']}) 原始分区键={orig_key}, 分区数={orig_cnt} → {t['V3判断']}")
else:
    print("  无")

# ==================== 2. 新加入的表 ====================
print("\n" + "=" * 70)
print(f"二、新加入的分区表（{len(newly_added)} 张）—— V3新增或原始标记为不分区")
print("=" * 70)

if newly_added:
    by_module = {}
    for en, info in newly_added.items():
        mod = info["模块"]
        if mod not in by_module:
            by_module[mod] = []
        by_module[mod].append(info)
    
    for mod in sorted(by_module.keys()):
        tables = by_module[mod]
        print(f"\n  [{mod}] {len(tables)}张:")
        for t in tables:
            print(f"    + {t['表中文名']}({t['表英文名']}) V3分区键={t['V3分区键']}, 分区数={t['V3分区数']} | 原始={t['原始判断']}")
else:
    print("  无")

# ==================== 3. 分区键/分区数变化 ====================
print("\n" + "=" * 70)
print(f"三、分区键/分区数变化的表（{len(both_partition)} 张）—— 两者都分区但参数不同")
print("=" * 70)

if both_partition:
    for en, info in both_partition.items():
        print(f"  ~ {info['表中文名']}({info['表英文名']}) [{info['模块']}]")
        print(f"    原始: 分区键={info['原始分区键']}, 分区数={info['原始分区数']}")
        print(f"    V3:   分区键={info['V3分区键']}, 分区数={info['V3分区数']}")
else:
    print("  无")

# ==================== 4. 保持一致的表 ====================
print("\n" + "=" * 70)
consistent = {k: v for k, v in v3_partition.items() if k in orig_partition and k not in both_partition}
print(f"四、保持一致的分区表（{len(consistent)} 张）")
print("=" * 70)
for en, info in sorted(consistent.items(), key=lambda x: x[1]["所属模块"]):
    print(f"  = {info['表中文名']}({info['表英文名']}) [{info['所属模块']}] 分区键={info['表分区键']}, 分区数={info['初始表分区数']}")

# ==================== 保存为JSON ====================
output = {
    "summary": {
        "原始分区表数": len(orig_partition),
        "V3分区表数": len(v3_partition),
        "被移除": len(removed),
        "新加入": len(newly_added),
        "分区键分区数变化": len(both_partition),
        "保持一致": len(consistent),
    },
    "removed": {k: v for k, v in removed.items()},
    "newly_added": {k: v for k, v in newly_added.items()},
    "changed": {k: v for k, v in both_partition.items()},
    "consistent": {k: v for k, v in consistent.items()},
}

with open(os.path.join(BASE_DIR, "compare_v3_vs_original.json"), "w", encoding="utf-8") as f:
    json.dump(output, f, ensure_ascii=False, indent=2)

print(f"\n对比结果已保存到: compare_v3_vs_original.json")