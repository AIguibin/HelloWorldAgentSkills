#!/usr/bin/env python3
"""
通用化 Excel 创建脚本，支持灵活的行/列格式配置。

支持通过 JSON 配置文件或直接调用函数来生成格式化的 Excel 文件。
"""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DEFAULT_FONT_FAMILIES = ["微软雅黑", "思源黑体", "WPS灵秀黑", "新宋体"]
FONT_NAME = "微软雅黑"
LIGHT_GREEN = "90EE90"


def hex_to_rgb(hex_color: str) -> str:
    return hex_color.lstrip("#").upper()


def create_border(color: str, style: str = "thin") -> Border:
    c = hex_to_rgb(color)
    side = Side(style=style, color=c)
    return Border(left=side, right=side, top=side, bottom=side)


def create_fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_to_rgb(color))


def create_font(
    name: str = FONT_NAME,
    size: int = 10,
    bold: bool = False,
    color: str = "000000",
) -> Font:
    return Font(name=name, size=size, bold=bold, color=hex_to_rgb(color))


def apply_row_format(ws, row_config: dict, row_idx: int, col_count: int):
    fill = create_fill(row_config["bg_color"]) if row_config.get("bg_color") else None
    font_cfg = row_config.get("font", {})
    font = create_font(
        name=font_cfg.get("name", FONT_NAME),
        size=font_cfg.get("size", 10),
        bold=font_cfg.get("bold", False),
        color=font_cfg.get("color", "000000"),
    )
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row_idx, column=col)
        if fill:
            cell.fill = fill
        cell.font = font
        if row_config.get("align"):
            cell.alignment = Alignment(
                horizontal=row_config["align"],
                vertical=row_config.get("vertical_align", "center"),
            )


def apply_borders(ws, start_row: int, end_row: int, col_count: int, border: Border):
    for row in range(start_row, end_row + 1):
        for col in range(1, col_count + 1):
            ws.cell(row=row, column=col).border = border


def set_column_widths(ws, col_count: int, width: int = 12):
    for col_idx in range(1, col_count + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def create_excel(config: dict, output_path: str) -> str:
    col_count = config.get("col_count", 10)
    row_range = config.get("row_range", 100)
    font_name = config.get("font_name", FONT_NAME)
    border_color = config.get("border_color", LIGHT_GREEN)
    vertical_align = config.get("vertical_align", "center")

    wb = Workbook()
    ws = wb.active
    ws.title = config.get("sheet_name", "Sheet1")

    border = create_border(border_color)

    for row_cfg in config.get("rows", []):
        row_idx = row_cfg["row"]
        start_col = row_cfg.get("start_col", 1)
        end_col = row_cfg.get("end_col", col_count)
        col_range = (start_col, end_col)
        align = row_cfg.get("align", None)

        if row_cfg.get("merge"):
            mc_start = row_cfg["merge"].get("start_col", start_col)
            mc_end = row_cfg["merge"].get("end_col", end_col)
            if isinstance(mc_start, str):
                mc_start = ord(mc_start.upper()) - ord("A") + 1
            if isinstance(mc_end, str):
                mc_end = ord(mc_end.upper()) - ord("A") + 1
            ws.merge_cells(
                start_row=row_idx, start_column=mc_start,
                end_row=row_idx, end_column=mc_end,
            )

        bg_color = row_cfg.get("bg_color")
        font_cfg = row_cfg.get("font", {})
        font = create_font(
            name=font_cfg.get("name", font_name),
            size=font_cfg.get("size", 10),
            bold=font_cfg.get("bold", False),
            color=font_cfg.get("color", "000000"),
        )
        fill = create_fill(bg_color) if bg_color else None

        for col in range(col_range[0], col_range[1] + 1):
            cell = ws.cell(row=row_idx, column=col)
            if fill:
                cell.fill = fill
            cell.font = font
            if align:
                cell.alignment = Alignment(horizontal=align, vertical=vertical_align)

        for cell_cfg in row_cfg.get("cells", []):
            col = cell_cfg["col"]
            c = ws.cell(row=row_idx, column=col)
            if "value" in cell_cfg:
                c.value = cell_cfg["value"]

    for row_cfg in config.get("default_rows", []):
        for row in range(row_cfg["start"], row_cfg["end"] + 1):
            font_cfg = row_cfg.get("font", {})
            font = create_font(
                name=font_cfg.get("name", font_name),
                size=font_cfg.get("size", 10),
                bold=font_cfg.get("bold", False),
                color=font_cfg.get("color", "000000"),
            )
            for col in range(1, col_count + 1):
                ws.cell(row=row, column=col).font = font

    apply_borders(ws, 1, row_range, col_count, border)
    set_column_widths(ws, col_count, config.get("column_width", 12))

    output = Path(output_path).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output))
    return str(output)


def main():
    if len(sys.argv) < 2:
        print("Usage: python create_excel.py <config.json> [output_path]")
        print("  config.json  - JSON configuration file")
        print("  output_path  - Output Excel file path (optional)")
        sys.exit(1)

    config_path = Path(sys.argv[1])
    if not config_path.exists():
        print(f"Error: Config file not found: {config_path}")
        sys.exit(1)

    with open(config_path, "r", encoding="utf-8") as f:
        config = json.load(f)

    output_path = sys.argv[2] if len(sys.argv) > 2 else config.get("output_path", "output.xlsx")
    result = create_excel(config, output_path)
    print(f"Excel file created: {result}")


if __name__ == "__main__":
    main()