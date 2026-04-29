#!/usr/bin/env python3
"""
Markdown任务文本转Excel待办清单

功能：
1. 解析Markdown格式的任务文本
2. 自动去重合并相似任务
3. 估算工作量（人天）
4. 生成规范化的Excel待办任务表

用法：
    python md2todo.py --input tasks.md --output todo.xlsx --owner "刘贵斌" --group "公共组"
"""

import argparse
import re
from datetime import datetime, timedelta
from typing import List, Dict, Tuple, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


class TaskParser:
    """任务解析器"""

    # 任务分类关键词映射
    CATEGORY_KEYWORDS = {
        "前端": ["前端", "UI", "样式", "页面", "组件", "按钮", "表单", "logo", "loading", "TAB"],
        "性能": ["性能", "优化", "SQL", "缓存", "压测", "启动", "编译", "耗时", "加载"],
        "安全": ["安全", "密码", "加密", "脱敏", "漏洞", "认证", "授权", "白名单"],
        "数据": ["数据", "字典", "表结构", "索引", "分区", "清理", "数据库"],
        "配置": ["配置", "yaml", "参数", "环境", "网关", "接口标准"],
        "设计": ["设计", "架构", "表结构"],
    }

    # 工作量估算规则 (任务类型 -> 复杂度关键词 -> 人天)
    EFFORT_ESTIMATES = {
        "前端": {
            "simple": ["样式", "UI", "logo", "对齐", "颜色"],
            "medium": ["组件", "表单", "校验", "搜索", "折叠"],
            "complex": ["框架", "性能", "优化", "重构"],
        },
        "性能": {
            "simple": ["参数", "配置"],
            "medium": ["SQL", "缓存", "索引"],
            "complex": ["架构", "压测", "熔断", "降级", "启动优化"],
        },
        "安全": {
            "simple": ["配置", "yaml", "加密"],
            "medium": ["修复", "漏洞"],
            "complex": ["组件", "脱敏", "认证改造"],
        },
        "开发": {
            "simple": ["配置", "参数", "对接"],
            "medium": ["功能", "开发", "接口", "联调"],
            "complex": ["拦截器", "同步机制", "权限", "工作流"],
        },
        "数据": {
            "simple": ["字典", "统一"],
            "medium": ["表结构", "索引", "分区"],
            "complex": ["清理策略", "评估", "实施"],
        },
        "配置": {
            "simple": ["yaml", "参数"],
            "medium": ["环境", "统一"],
            "complex": ["网关", "接口标准"],
        },
        "设计": {
            "simple": ["文档"],
            "medium": ["表结构"],
            "complex": ["架构"],
        },
    }

    # 复杂度对应人天
    COMPLEXITY_DAYS = {"simple": 3, "medium": 5, "complex": 8}

    def __init__(self):
        self.tasks: List[Dict] = []

    def parse_text(self, text: str) -> List[Dict]:
        """解析文本提取任务"""
        lines = text.strip().split("\n")

        for line in lines:
            line = line.strip()
            if not line or line.startswith("#"):
                continue

            task = self._parse_line(line)
            if task:
                self.tasks.append(task)

        # 去重
        self.tasks = self._deduplicate_tasks(self.tasks)

        return self.tasks

    def _parse_line(self, line: str) -> Optional[Dict]:
        """解析单行文本"""
        # 匹配模式：负责人--任务描述--状态
        pattern1 = r"^(\w+)\s*[-—]+\s*(.+?)(?:\s*[-—]+\s*(.+))?$"

        # 匹配模式：序号. 负责人---任务
        pattern2 = r"^\d+\s*[-.]*\s*(\w+)\s*[-—]+(.+)$"

        match = re.match(pattern1, line) or re.match(pattern2, line)

        if match:
            owner = match.group(1).strip()
            desc = match.group(2).strip()
            status_str = match.group(3).strip() if len(match.groups()) > 2 and match.group(3) else ""

            # 解析状态
            status = self._parse_status(status_str)

            # 判断分类
            category = self._categorize_task(desc)

            # 判断优先级
            priority = self._determine_priority(desc, category)

            # 估算工作量
            effort_days = self._estimate_effort(desc, category)

            # 提取关联人员
            related = self._extract_related(desc)

            return {
                "owner": owner,
                "description": desc,
                "status": status,
                "category": category,
                "priority": priority,
                "effort_days": effort_days,
                "related": related,
            }

        return None

    def _parse_status(self, status_str: str) -> str:
        """解析任务状态"""
        status_str = status_str.lower()

        if any(kw in status_str for kw in ["已完成", "完成", "done", "结束"]):
            return "已完成"
        elif any(kw in status_str for kw in ["进行中", "开发中", "排期中"]):
            return "进行中"
        else:
            return "未开始"

    def _categorize_task(self, desc: str) -> str:
        """判断任务分类"""
        desc_lower = desc.lower()

        for category, keywords in self.CATEGORY_KEYWORDS.items():
            if any(kw in desc_lower for kw in keywords):
                return category

        return "开发"

    def _determine_priority(self, desc: str, category: str) -> str:
        """判断优先级"""
        desc_lower = desc.lower()

        # 高优先级关键词
        high_keywords = ["性能", "安全", "漏洞", "核心", "重要", "紧急", "高", "线上", "生产"]
        if any(kw in desc_lower for kw in high_keywords):
            return "高"

        # 低优先级关键词
        low_keywords = ["UI", "样式", "logo", "对齐", "颜色", "优化体验"]
        if any(kw in desc_lower for kw in low_keywords):
            return "低"

        return "中"

    def _estimate_effort(self, desc: str, category: str) -> int:
        """估算工作量（人天）"""
        desc_lower = desc.lower()

        # 获取该分类的估算规则
        rules = self.EFFORT_ESTIMATES.get(category, self.EFFORT_ESTIMATES["开发"])

        # 检查复杂度关键词
        for complexity, keywords in rules.items():
            if any(kw in desc_lower for kw in keywords):
                return self.COMPLEXITY_DAYS[complexity]

        # 默认中等复杂度
        return 5

    def _extract_related(self, desc: str) -> str:
        """提取关联人员/组"""
        # 匹配"与xxx沟通"、"对接xxx"等模式
        patterns = [
            r"与(\w+)沟通",
            r"对接(\w+)",
            r"协调(\w+)",
        ]

        related = []
        for pattern in patterns:
            matches = re.findall(pattern, desc)
            related.extend(matches)

        return "、".join(related) if related else ""

    def _deduplicate_tasks(self, tasks: List[Dict]) -> List[Dict]:
        """任务去重"""
        unique_tasks = {}

        for task in tasks:
            # 使用负责人+任务描述前20字作为去重键
            key = f"{task['owner']}_{task['description'][:20]}"

            if key in unique_tasks:
                # 保留更完整的描述
                if len(task["description"]) > len(unique_tasks[key]["description"]):
                    unique_tasks[key] = task
            else:
                unique_tasks[key] = task

        return list(unique_tasks.values())


class ExcelGenerator:
    """Excel生成器"""

    def __init__(self, owner: str = "刘贵斌", group: str = "公共组", start_date: Optional[str] = None):
        self.owner = owner
        self.group = group
        self.start_date = datetime.strptime(start_date, "%Y-%m-%d") if start_date else datetime.now()
        self.register_date = self.start_date.strftime("%Y-%m-%d")

    def generate(self, tasks: List[Dict], output_path: str):
        """生成Excel文件"""
        wb = Workbook()
        ws = wb.active
        ws.title = "待办任务清单"

        # 写入表头
        headers = [
            "登记时间", "登记人", "事项分类", "小组", "事项描述",
            "优先级", "参与人员", "进度", "计划开始时间", "计划完成时间",
            "负责人", "涉及关联人员/组", "进度说明"
        ]
        ws.append(headers)

        # 设置表头样式
        self._set_header_style(ws)

        # 写入任务数据
        for task in tasks:
            row = self._create_row(task)
            ws.append(row)

        # 设置列宽
        self._set_column_widths(ws)

        # 设置边框和对齐
        self._set_cell_styles(ws)

        # 冻结首行
        ws.freeze_panes = "A2"

        # 保存文件
        wb.save(output_path)
        print(f"✅ Excel文件已生成: {output_path}")
        print(f"📊 共整理 {len(tasks)} 个待办任务")

    def _create_row(self, task: Dict) -> List:
        """创建Excel行数据"""
        # 计算计划时间
        effort_days = task["effort_days"]

        # 根据状态调整开始时间偏移
        if task["status"] == "进行中":
            start_offset = 0
        elif task["priority"] == "高":
            start_offset = 3  # 高优先级延后3天开始
        else:
            start_offset = 0

        start_date = self.start_date + timedelta(days=start_offset)
        end_date = start_date + timedelta(days=effort_days)

        return [
            self.register_date,           # 登记时间
            self.owner,                   # 登记人
            task["category"],             # 事项分类
            self.group,                   # 小组
            task["description"],          # 事项描述
            task["priority"],             # 优先级
            "",                           # 参与人员
            task["status"],               # 进度
            start_date.strftime("%Y-%m-%d"),  # 计划开始时间
            end_date.strftime("%Y-%m-%d"),    # 计划完成时间
            task["owner"],                # 负责人
            task["related"],              # 涉及关联人员/组
            "",                           # 进度说明
        ]

    def _set_header_style(self, ws):
        """设置表头样式"""
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _set_column_widths(self, ws):
        """设置列宽"""
        widths = [12, 10, 10, 10, 60, 8, 12, 10, 14, 14, 12, 20, 30]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = width

    def _set_cell_styles(self, ws):
        """设置单元格样式"""
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=13):
            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center", wrap_text=True)


def main():
    parser = argparse.ArgumentParser(description="Markdown任务文本转Excel待办清单")
    parser.add_argument("--input", "-i", required=True, help="输入Markdown文件路径")
    parser.add_argument("--output", "-o", required=True, help="输出Excel文件路径")
    parser.add_argument("--owner", default="刘贵斌", help="登记人姓名")
    parser.add_argument("--group", default="公共组", help="所属小组")
    parser.add_argument("--start-date", help="计划开始日期 (YYYY-MM-DD)")

    args = parser.parse_args()

    # 读取输入文件
    with open(args.input, "r", encoding="utf-8") as f:
        text = f.read()

    # 解析任务
    parser = TaskParser()
    tasks = parser.parse_text(text)

    if not tasks:
        print("⚠️ 未识别到任何任务，请检查输入文件格式")
        return

    # 生成Excel
    generator = ExcelGenerator(
        owner=args.owner,
        group=args.group,
        start_date=args.start_date
    )
    generator.generate(tasks, args.output)


if __name__ == "__main__":
    main()
