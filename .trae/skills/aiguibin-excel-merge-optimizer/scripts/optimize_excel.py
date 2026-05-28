#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import argparse
from openpyxl import load_workbook


def find_rows_to_delete(ws):
    """
    找出所有需要删除的空白行和重复标题行
    
    Args:
        ws: 工作表对象
    
    Returns:
        tuple: (需要删除的行列表, 空白行数量, 重复标题行数量)
    """
    # 获取第一行标题
    first_row_header = []
    for col in range(1, ws.max_column + 1):
        cell_value = ws.cell(row=1, column=col).value
        first_row_header.append(str(cell_value) if cell_value else "")
    
    # 找出所有空白行和重复标题行
    empty_rows = []
    duplicate_header_rows = []
    
    for row_idx in range(2, ws.max_row + 1):
        # 检查是否为空白行
        is_empty = True
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=row_idx, column=col).value
            if cell_value is not None and str(cell_value).strip() != "":
                is_empty = False
                break
        
        if is_empty:
            empty_rows.append(row_idx)
            continue
        
        # 检查是否为重复标题行
        current_row_values = []
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=row_idx, column=col).value
            current_row_values.append(str(cell_value) if cell_value else "")
        
        if current_row_values == first_row_header:
            duplicate_header_rows.append(row_idx)
    
    # 合并需要删除的行并排序
    rows_to_delete = sorted(set(empty_rows + duplicate_header_rows))
    
    return rows_to_delete, len(empty_rows), len(duplicate_header_rows)


def optimize_excel_file(file_path):
    """
    删除Excel文件中的空白行和重复标题行
    
    Args:
        file_path: Excel文件路径
    
    Returns:
        tuple: (是否成功, 原始行数, 优化后行数, 删除行数)
    """
    print(f"正在处理文件: {file_path}")
    
    try:
        # 加载工作簿
        wb = load_workbook(file_path, data_only=False)
        ws = wb.active
        
        print(f"工作表名称: {ws.title}")
        original_rows = ws.max_row
        print(f"原始行数: {original_rows}")
        
        # 找出需要删除的行
        rows_to_delete, empty_count, duplicate_count = find_rows_to_delete(ws)
        
        print(f"\n找到 {empty_count} 个空白行")
        print(f"找到 {duplicate_count} 个重复标题行")
        print(f"总共需要删除 {len(rows_to_delete)} 行")
        
        # 从后往前删除行（避免行号变化）
        deleted_count = 0
        for row_idx in reversed(rows_to_delete):
            ws.delete_rows(row_idx)
            deleted_count += 1
            if deleted_count % 10 == 0:
                print(f"  已删除 {deleted_count}/{len(rows_to_delete)} 行...")
        
        print(f"\n删除完成!")
        optimized_rows = ws.max_row
        print(f"删除后行数: {optimized_rows}")
        
        # 保存文件
        print(f"\n保存文件: {file_path}")
        wb.save(file_path)
        wb.close()
        
        print("优化完成!")
        
        return True, original_rows, optimized_rows, deleted_count
        
    except Exception as e:
        print(f"错误: 处理文件时发生异常 - {str(e)}")
        return False, 0, 0, 0


def verify_optimized_file(file_path):
    """
    验证优化后的文件
    
    Args:
        file_path: Excel文件路径
    
    Returns:
        tuple: (是否验证通过, 重复标题行数, 空白行数)
    """
    print(f"\n验证优化后的文件...")
    
    try:
        wb = load_workbook(file_path, data_only=False)
        ws = wb.active
        
        print(f"工作表名称: {ws.title}")
        print(f"总行数: {ws.max_row}")
        print(f"总列数: {ws.max_column}")
        
        # 检查是否还有重复标题
        first_row_header = []
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=1, column=col).value
            first_row_header.append(str(cell_value) if cell_value else "")
        
        duplicate_count = 0
        for row_idx in range(2, ws.max_row + 1):
            current_row_values = []
            for col in range(1, ws.max_column + 1):
                cell_value = ws.cell(row=row_idx, column=col).value
                current_row_values.append(str(cell_value) if cell_value else "")
            
            if current_row_values == first_row_header:
                duplicate_count += 1
        
        if duplicate_count == 0:
            print("✓ 验证通过: 没有重复标题行")
        else:
            print(f"✗ 验证失败: 仍有 {duplicate_count} 个重复标题行")
        
        # 检查是否还有空白行
        empty_count = 0
        for row_idx in range(2, ws.max_row + 1):
            is_empty = True
            for col in range(1, ws.max_column + 1):
                cell_value = ws.cell(row=row_idx, column=col).value
                if cell_value is not None and str(cell_value).strip() != "":
                    is_empty = False
                    break
            if is_empty:
                empty_count += 1
        
        if empty_count == 0:
            print("✓ 验证通过: 没有空白行")
        else:
            print(f"✗ 验证失败: 仍有 {empty_count} 个空白行")
        
        # 显示前几行数据
        print("\n数据预览（前5行）:")
        for row_idx in range(1, min(6, ws.max_row + 1)):
            row_data = []
            for col in range(1, min(6, ws.max_column + 1)):
                cell_value = ws.cell(row=row_idx, column=col).value
                row_data.append(str(cell_value) if cell_value else "")
            print(f"  行{row_idx}: {' | '.join(row_data)}")
        
        # 显示最后几行数据
        print("\n数据预览（最后5行）:")
        for row_idx in range(max(1, ws.max_row - 4), ws.max_row + 1):
            row_data = []
            for col in range(1, min(6, ws.max_column + 1)):
                cell_value = ws.cell(row=row_idx, column=col).value
                row_data.append(str(cell_value) if cell_value else "")
            print(f"  行{row_idx}: {' | '.join(row_data)}")
        
        wb.close()
        
        is_valid = duplicate_count == 0 and empty_count == 0
        return is_valid, duplicate_count, empty_count
        
    except Exception as e:
        print(f"验证失败: {str(e)}")
        return False, -1, -1


def main():
    parser = argparse.ArgumentParser(description='优化Excel文件格式，删除空白行和重复标题行')
    parser.add_argument('--file', required=True, help='Excel文件路径')
    
    args = parser.parse_args()
    
    # 执行优化
    success, original_rows, optimized_rows, deleted_count = optimize_excel_file(args.file)
    
    if success:
        # 验证结果
        is_valid, duplicate_count, empty_count = verify_optimized_file(args.file)
        
        print(f"\n{'='*60}")
        print("优化报告:")
        print(f"{'='*60}")
        print(f"原始行数: {original_rows}")
        print(f"优化后行数: {optimized_rows}")
        print(f"删除行数: {deleted_count}")
        print(f"验证结果: {'通过' if is_valid else '失败'}")
        print(f"{'='*60}")
        
        return 0 if is_valid else 1
    else:
        print(f"\n✗ 优化失败!")
        return 1


if __name__ == "__main__":
    exit(main())
