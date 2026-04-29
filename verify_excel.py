#!/usr/bin/env python
# -*- coding: utf-8 -*-
import openpyxl

# 文件路径
file_path = r"e:\AutoOffice\zzzzzzzz-common-inbox\todo-inbox\公共组工作计划.xlsx"

# 加载工作簿
print("验证Excel文件...")
wb = openpyxl.load_workbook(file_path)
ws = wb['Sheet1']

# 状态列是F列（第6列）
status_col = 6

# 统计各种状态的任务数量
status_count = {}
completed_rows = []

for row_idx in range(2, ws.max_row + 1):
    status_value = ws.cell(row=row_idx, column=status_col).value
    if status_value:
        status_count[status_value] = status_count.get(status_value, 0) + 1
        if status_value == "已完成":
            completed_rows.append(row_idx)

print(f"\n文件统计信息:")
print(f"  总行数: {ws.max_row}")
print(f"  数据行数: {ws.max_row - 1}")
print(f"\n各状态任务数量:")
for status, count in sorted(status_count.items()):
    print(f"  {status}: {count} 行")

if completed_rows:
    print(f"\n警告: 仍有 {len(completed_rows)} 行'已完成'状态的任务!")
    for row_idx in completed_rows[:5]:
        wbs = ws.cell(row=row_idx, column=1).value
        name = ws.cell(row=row_idx, column=2).value
        print(f"  行{row_idx}: WBS={wbs}, 名称={name}")
else:
    print(f"\n✓ 验证通过: 所有'已完成'状态的任务已被成功删除!")
