#!/usr/bin/env python
# -*- coding: utf-8 -*-
import openpyxl
import shutil

# 文件路径
file_path = r"e:\AutoOffice\zzzzzzzz-common-inbox\todo-inbox\公共组工作计划.xlsx"
backup_path = r"e:\AutoOffice\zzzzzzzz-common-inbox\todo-inbox\公共组工作计划_backup.xlsx"

# 创建备份
print("创建备份...")
shutil.copy2(file_path, backup_path)
print(f"备份已创建: {backup_path}")

# 加载工作簿
print("加载Excel文件...")
wb = openpyxl.load_workbook(file_path)
ws = wb['Sheet1']

# 状态列是F列（第6列）
status_col = 6

# 收集需要删除的行（从后往前删）
rows_to_delete = []
print("扫描所有行，查找状态为'已完成'的任务...")

for row_idx in range(2, ws.max_row + 1):  # 从第2行开始（跳过表头）
    status_value = ws.cell(row=row_idx, column=status_col).value
    if status_value == "已完成":
        rows_to_delete.append(row_idx)

print(f"找到 {len(rows_to_delete)} 行状态为'已完成'的任务")

# 打印前20个要删除的任务
print("\n前20个要删除的任务:")
for i, row_idx in enumerate(rows_to_delete[:20]):
    wbs = ws.cell(row=row_idx, column=1).value
    name = ws.cell(row=row_idx, column=2).value
    print(f"  行{row_idx}: WBS={wbs}, 名称={name[:40] if name else ''}")

if len(rows_to_delete) > 20:
    print(f"  ... 还有 {len(rows_to_delete) - 20} 行")

# 从后往前删除行
print(f"\n开始删除 {len(rows_to_delete)} 行...")
for row_idx in reversed(rows_to_delete):
    ws.delete_rows(row_idx)

print(f"已删除 {len(rows_to_delete)} 行")

# 保存文件
print("保存文件...")
wb.save(file_path)
print(f"文件已保存: {file_path}")
print(f"\n操作完成！共删除 {len(rows_to_delete)} 行'已完成'状态的任务")
