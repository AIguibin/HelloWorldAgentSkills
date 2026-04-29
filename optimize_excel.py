#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from copy import copy


def find_rows_to_delete(ws):
    """
    找出所有需要删除的空白行和重复标题行
    """
    # 获取第一行标题
    first_row_header = []
    for col in range(1, ws.max_column + 1):
        cell_value = ws.cell(row=1, column=col).value
        first_row_header.append(str(cell_value) if cell_value else "")
    
    # 找出所有空白行和重复标题行
    empty_rows = []
    duplicate_header_rows = []
    
    for row_idx in range(2, ws.max_row + 1):
        # 检查是否为空白行
        is_empty = True
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=row_idx, column=col).value
            if cell_value is not None and str(cell_value).strip() != "":
                is_empty = False
                break
        
        if is_empty:
            empty_rows.append(row_idx)
            continue
        
        # 检查是否为重复标题行
        current_row_values = []
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=row_idx, column=col).value
            current_row_values.append(str(cell_value) if cell_value else "")
        
        if current_row_values == first_row_header:
            duplicate_header_rows.append(row_idx)
    
    # 合并需要删除的行并排序
    rows_to_delete = sorted(set(empty_rows + duplicate_header_rows))
    
    return rows_to_delete, len(empty_rows), len(duplicate_header_rows)


def delete_rows_from_excel(file_path):
    """
    删除Excel文件中的空白行和重复标题行
    """
    print(f"正在处理文件: {file_path}")
    
    # 加载工作簿
    wb = load_workbook(file_path, data_only=False)
    ws = wb.active
    
    print(f"工作表名称: {ws.title}")
    print(f"原始行数: {ws.max_row}")
    
    # 找出需要删除的行
    rows_to_delete, empty_count, duplicate_count = find_rows_to_delete(ws)
    
    print(f"\n找到 {empty_count} 个空白行")
    print(f"找到 {duplicate_count} 个重复标题行")
    print(f"总共需要删除 {len(rows_to_delete)} 行")
    
    # 从后往前删除行（避免行号变化）
    deleted_count = 0
    for row_idx in reversed(rows_to_delete):
        ws.delete_rows(row_idx)
        deleted_count += 1
        if deleted_count % 10 == 0:
            print(f"  已删除 {deleted_count}/{len(rows_to_delete)} 行...")
    
    print(f"\n删除完成!")
    print(f"删除后行数: {ws.max_row}")
    
    # 保存文件
    print(f"\n保存文件: {file_path}")
    wb.save(file_path)
    wb.close()
    
    print("优化完成!")
    
    return ws.max_row


def verify_optimized_file(file_path):
    """
    验证优化后的文件
    """
    print(f"\n验证优化后的文件...")
    
    wb = load_workbook(file_path, data_only=False)
    ws = wb.active
    
    print(f"工作表名称: {ws.title}")
    print(f"总行数: {ws.max_row}")
    print(f"总列数: {ws.max_column}")
    
    # 检查是否还有重复标题
    first_row_header = []
    for col in range(1, ws.max_column + 1):
        cell_value = ws.cell(row=1, column=col).value
        first_row_header.append(str(cell_value) if cell_value else "")
    
    duplicate_count = 0
    for row_idx in range(2, ws.max_row + 1):
        current_row_values = []
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=row_idx, column=col).value
            current_row_values.append(str(cell_value) if cell_value else "")
        
        if current_row_values == first_row_header:
            duplicate_count += 1
    
    if duplicate_count == 0:
        print("✓ 验证通过: 没有重复标题行")
    else:
        print(f"✗ 验证失败: 仍有 {duplicate_count} 个重复标题行")
    
    # 检查是否还有空白行
    empty_count = 0
    for row_idx in range(2, ws.max_row + 1):
        is_empty = True
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=row_idx, column=col).value
            if cell_value is not None and str(cell_value).strip() != "":
                is_empty = False
                break
        if is_empty:
            empty_count += 1
    
    if empty_count == 0:
        print("✓ 验证通过: 没有空白行")
    else:
        print(f"✗ 验证失败: 仍有 {empty_count} 个空白行")
    
    # 显示前几行数据
    print("\n数据预览（前5行）:")
    for row_idx in range(1, min(6, ws.max_row + 1)):
        row_data = []
        for col in range(1, min(6, ws.max_column + 1)):
            cell_value = ws.cell(row=row_idx, column=col).value
            row_data.append(str(cell_value) if cell_value else "")
        print(f"  行{row_idx}: {' | '.join(row_data)}")
    
    # 显示最后几行数据
    print("\n数据预览（最后5行）:")
    for row_idx in range(max(1, ws.max_row - 4), ws.max_row + 1):
        row_data = []
        for col in range(1, min(6, ws.max_column + 1)):
            cell_value = ws.cell(row=row_idx, column=col).value
            row_data.append(str(cell_value) if cell_value else "")
        print(f"  行{row_idx}: {' | '.join(row_data)}")
    
    wb.close()
    
    return duplicate_count == 0 and empty_count == 0


if __name__ == "__main__":
    file_path = r"e:\WorkSpace\HelloWorldAgentSkills\aiguibin-common-excel\湖北农信新信贷数据库设计文档\新信贷_数据库设计_全量信贷v1.0.0.xlsx"
    
    # 删除多余的行
    delete_rows_from_excel(file_path)
    
    # 验证结果
    verify_optimized_file(file_path)
