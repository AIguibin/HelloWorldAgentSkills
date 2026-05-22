from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Sheet1"

font_name = "微软雅黑"
purple_fill = PatternFill("solid", fgColor="7030A0")
green_fill = PatternFill("solid", fgColor="00B050")
light_green_border = Border(
    left=Side(style="thin", color="90EE90"),
    right=Side(style="thin", color="90EE90"),
    top=Side(style="thin", color="90EE90"),
    bottom=Side(style="thin", color="90EE90"),
)

ws.merge_cells("A1:J1")
cell_a1 = ws["A1"]
cell_a1.font = Font(name=font_name, size=16, bold=True, color="FFFFFF")
cell_a1.fill = purple_fill
cell_a1.alignment = Alignment(horizontal="center", vertical="center")

for col in range(1, 11):
    cell = ws.cell(row=2, column=col)
    cell.fill = green_fill
    cell.font = Font(name=font_name, size=12, bold=True, color="FFFFFF")

ws["A2"] = "序号"
ws["J2"] = "备注"

for row in range(1, 101):
    for col in range(1, 11):
        cell = ws.cell(row=row, column=col)
        cell.border = light_green_border

for row in range(3, 101):
    for col in range(1, 11):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=font_name, size=10)

for col_idx in range(1, 11):
    col_letter = get_column_letter(col_idx)
    ws.column_dimensions[col_letter].width = 12

output_path = r"e:\WorkSpace\HelloWorldAgentSkills\aiguibin-common-create\formatted_excel.xlsx"
wb.save(output_path)
print(f"Excel file saved to: {output_path}")