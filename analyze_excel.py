#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from collections import defaultdict


def analyze_excel_file(file_path):
    """
    分析Excel文件，找出所有空白行和重复标题行
    """
    wb = load_workbook(file_path, data_only=False)
    ws = wb.active
    
    print(f"工作表名称: {ws.title}")
    print(f"总行数: {ws.max_row}")
    print(f"总列数: {ws.max_column}")
    
    # 获取第一行标题
    first_row_header = []
    for col in range(1, ws.max_column + 1):
        cell_value = ws.cell(row=1, column=col).value
        first_row_header.append(str(cell_value) if cell_value else "")
    
    print(f"\n第一行标题: {first_row_header[:5]}...")
    
    # 找出所有空白行和重复标题行
    empty_rows = []
    duplicate_header_rows = []
    
    for row_idx in range(2, ws.max_row + 1):
        # 检查是否为空白行
        is_empty = True
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=row_idx, column=col).value
            if cell_value is not None and str(cell_value).strip() != "":
                is_empty = False
                break
        
        if is_empty:
            empty_rows.append(row_idx)
            continue
        
        # 检查是否为重复标题行
        current_row_values = []
        for col in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=row_idx, column=col).value
            current_row_values.append(str(cell_value) if cell_value else "")
        
        if current_row_values == first_row_header:
            duplicate_header_rows.append(row_idx)
    
    print(f"\n找到 {len(empty_rows)} 个空白行:")
    if len(empty_rows) <= 20:
        print(f"  行号: {empty_rows}")
    else:
        print(f"  前20个: {empty_rows[:20]}")
        print(f"  ... 还有 {len(empty_rows) - 20} 个")
    
    print(f"\n找到 {len(duplicate_header_rows)} 个重复标题行:")
    if len(duplicate_header_rows) <= 20:
        print(f"  行号: {duplicate_header_rows}")
    else:
        print(f"  前20个: {duplicate_header_rows[:20]}")
        print(f"  ... 还有 {len(duplicate_header_rows) - 20} 个")
    
    # 合并需要删除的行
    rows_to_delete = sorted(set(empty_rows + duplicate_header_rows))
    print(f"\n总共需要删除 {len(rows_to_delete)} 行")
    
    wb.close()
    
    return rows_to_delete


if __name__ == "__main__":
    file_path = r"e:\WorkSpace\HelloWorldAgentSkills\aiguibin-common-excel\湖北农信新信贷数据库设计文档\新信贷_数据库设计_全量信贷v1.0.0.xlsx"
    analyze_excel_file(file_path)
