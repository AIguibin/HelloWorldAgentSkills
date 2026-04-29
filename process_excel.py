#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel处理脚本：删除状态为"已完成"的任务行
"""

import openpyxl
from openpyxl import load_workbook
import os

def remove_completed_tasks(file_path):
    """
    删除Excel中状态为"已完成"的任务行
    """
    print(f"正在处理文件: {file_path}")
    
    # 加载工作簿
    wb = load_workbook(file_path)
    ws = wb['Sheet1']
    
    # 获取状态列的索引 (F列 = 第6列)
    status_col = 6  # F列
    
    # 收集需要删除的行号（从后往前删，避免索引变化问题）
    rows_to_delete = []
    
    # 遍历所有行，从第2行开始（第1行是表头）
    for row_idx in range(2, ws.max_row + 1):
        status_cell = ws.cell(row=row_idx, column=status_col)
        status_value = status_cell.value
        
        if status_value == "已完成":
            rows_to_delete.append(row_idx)
    
    print(f"找到 {len(rows_to_delete)} 行状态为'已完成'的任务")
    
    # 打印前10个要删除的任务信息
    print("\n前10个要删除的任务:")
    for i, row_idx in enumerate(rows_to_delete[:10]):
        wbs = ws.cell(row=row_idx, column=1).value  # WBS编号
        name = ws.cell(row=row_idx, column=2).value  # 计划名称
        print(f"  行{row_idx}: WBS={wbs}, 名称={name}")
    
    if len(rows_to_delete) > 10:
        print(f"  ... 还有 {len(rows_to_delete) - 10} 行")
    
    # 从后往前删除行（这样删除不会影响前面行的索引）
    print(f"\n开始删除 {len(rows_to_delete)} 行...")
    for row_idx in reversed(rows_to_delete):
        ws.delete_rows(row_idx)
    
    print(f"已删除 {len(rows_to_delete)} 行")
    
    # 保存文件
    wb.save(file_path)
    print(f"文件已保存: {file_path}")
    
    return len(rows_to_delete)

if __name__ == "__main__":
    file_path = r"e:\AutoOffice\zzzzzzzz-common-inbox\todo-inbox\公共组工作计划.xlsx"
    
    # 创建备份
    backup_path = file_path.replace('.xlsx', '_backup.xlsx')
    print(f"创建备份: {backup_path}")
    
    import shutil
    shutil.copy2(file_path, backup_path)
    
    # 处理文件
    deleted_count = remove_completed_tasks(file_path)
    print(f"\n操作完成！共删除 {deleted_count} 行'已完成'状态的任务")
