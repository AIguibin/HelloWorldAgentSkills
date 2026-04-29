#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os
import re
import argparse
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
from copy import copy


def natural_sort_key(s):
    """
    生成自然排序的键，用于按文件名中的数字进行排序
    """
    return [int(text) if text.isdigit() else text.lower() for text in re.split('([0-9]+)', str(s))]


def copy_cell_style(source_cell, target_cell):
    """
    复制单元格的所有样式属性
    """
    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.border = copy(source_cell.border)
        target_cell.fill = copy(source_cell.fill)
        target_cell.number_format = copy(source_cell.number_format)
        target_cell.protection = copy(source_cell.protection)
        target_cell.alignment = copy(source_cell.alignment)


def copy_row_dimensions(source_sheet, target_sheet, source_row, target_row):
    """
    复制行高
    """
    if source_row in source_sheet.row_dimensions:
        source_dim = source_sheet.row_dimensions[source_row]
        target_sheet.row_dimensions[target_row].height = source_dim.height


def copy_column_dimensions(source_sheet, target_sheet):
    """
    复制列宽
    """
    for col_letter, col_dim in source_sheet.column_dimensions.items():
        if col_letter not in target_sheet.column_dimensions:
            target_sheet.column_dimensions[col_letter].width = col_dim.width
        else:
            if col_dim.width and col_dim.width > target_sheet.column_dimensions[col_letter].width:
                target_sheet.column_dimensions[col_letter].width = col_dim.width


def copy_merged_cells(source_sheet, target_sheet, row_offset):
    """
    复制合并单元格，并调整行号偏移
    """
    for merged_range in source_sheet.merged_cells.ranges:
        new_range = f"{get_column_letter(merged_range.min_col)}{merged_range.min_row + row_offset}:{get_column_letter(merged_range.max_col)}{merged_range.max_row + row_offset}"
        target_sheet.merge_cells(new_range)


def merge_excel_files(folder_path, sheet_name, output_file):
    """
    合并文件夹中所有Excel文件的指定工作表
    
    Args:
        folder_path: 源文件夹路径
        sheet_name: 要合并的工作表名称
        output_file: 输出文件路径
    
    Returns:
        tuple: (处理文件数量, 总行数, 处理的文件列表)
    """
    folder = Path(folder_path)
    
    # 获取所有xlsx文件（排除临时文件和输出文件）
    excel_files = []
    for file in folder.glob("*.xlsx"):
        if not file.name.startswith("~$") and file.name != Path(output_file).name:
            excel_files.append(file)
    
    # 按自然顺序排序
    excel_files.sort(key=natural_sort_key)
    
    print(f"找到 {len(excel_files)} 个Excel文件:")
    for idx, file in enumerate(excel_files, 1):
        print(f"  {idx}. {file.name}")
    
    # 创建新的工作簿
    new_wb = Workbook()
    new_ws = new_wb.active
    new_ws.title = sheet_name
    
    current_row = 1
    processed_files = []
    total_data_rows = 0
    
    for file_idx, excel_file in enumerate(excel_files, 1):
        print(f"\n处理文件 {file_idx}/{len(excel_files)}: {excel_file.name}")
        
        try:
            # 加载工作簿
            wb = load_workbook(excel_file, data_only=False)
            
            # 查找指定工作表
            target_sheet = None
            for sname in wb.sheetnames:
                if sname == sheet_name:
                    target_sheet = wb[sname]
                    break
            
            if target_sheet is None:
                print(f"  警告: 未找到'{sheet_name}'工作表，跳过此文件")
                wb.close()
                continue
            
            # 获取数据范围
            max_row = target_sheet.max_row
            max_col = target_sheet.max_column
            
            if max_row == 0 or max_col == 0:
                print(f"  警告: '{sheet_name}'工作表为空，跳过此文件")
                wb.close()
                continue
            
            print(f"  找到'{sheet_name}'工作表，数据范围: {max_row}行 x {max_col}列")
            
            # 复制列宽（只在第一次时复制）
            if file_idx == 1:
                copy_column_dimensions(target_sheet, new_ws)
            
            # 复制合并单元格
            row_offset = current_row - 1
            copy_merged_cells(target_sheet, new_ws, row_offset)
            
            # 复制数据和样式
            for row_idx in range(1, max_row + 1):
                for col_idx in range(1, max_col + 1):
                    source_cell = target_sheet.cell(row=row_idx, column=col_idx)
                    target_cell = new_ws.cell(row=current_row, column=col_idx)
                    
                    # 复制值
                    target_cell.value = source_cell.value
                    
                    # 复制样式
                    copy_cell_style(source_cell, target_cell)
                
                # 复制行高
                copy_row_dimensions(target_sheet, new_ws, row_idx, current_row)
                
                current_row += 1
            
            processed_files.append(excel_file.name)
            total_data_rows += max_row
            print(f"  成功合并 {max_row} 行数据")
            
            wb.close()
            
        except Exception as e:
            print(f"  错误: 处理文件时发生异常 - {str(e)}")
            continue
    
    # 保存新文件
    print(f"\n保存合并后的文件: {output_file}")
    new_wb.save(output_file)
    
    print(f"\n合并完成!")
    print(f"成功处理 {len(processed_files)} 个文件")
    print(f"总数据行数: {total_data_rows}")
    print(f"输出文件: {output_file}")
    
    return len(processed_files), total_data_rows, processed_files


def main():
    parser = argparse.ArgumentParser(description='合并多个Excel文件的指定工作表')
    parser.add_argument('--folder', required=True, help='源文件夹路径')
    parser.add_argument('--sheet', default='目录', help='工作表名称（默认：目录）')
    parser.add_argument('--output', required=True, help='输出文件路径')
    
    args = parser.parse_args()
    
    # 执行合并
    processed_count, total_rows, processed_files = merge_excel_files(
        args.folder, 
        args.sheet, 
        args.output
    )
    
    # 返回处理结果
    if processed_count > 0:
        print(f"\n✓ 合并成功!")
        print(f"  处理文件数: {processed_count}")
        print(f"  总数据行数: {total_rows}")
        print(f"  输出文件: {args.output}")
        return 0
    else:
        print(f"\n✗ 合并失败!")
        return 1


if __name__ == "__main__":
    exit(main())
