"""提取湖北农信清理归档表清单.xlsx 的完整数据"""
import json
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

try:
    from openpyxl import load_workbook
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    from openpyxl import load_workbook

file_path = r"e:\WorkSpace\HelloWorldAgentSkills\aiguibin-common-excel\DD-湖北农信基础配置检核\湖北农信清理归档表清单.xlsx"
wb = load_workbook(file_path, data_only=True)
ws = wb["目录"]

rows = []
headers = [cell.value for cell in ws[1]]
print("表头:", headers)
print(f"总行数(含表头): {ws.max_row}")
print(f"总列数: {ws.max_column}")
print("=" * 80)

for row in ws.iter_rows(min_row=2, values_only=True):
    if row[0] is None:
        continue
    row_data = {
        "序号": row[0],
        "所属数据库": row[1],
        "所属模块": row[2],
        "责任人": row[3],
        "表名_英文": row[4],
        "表名_中文": row[5],
        "是否数据清理": row[6],
        "清理策略": row[7],
        "归档策略": row[8],
        "备注说明": row[9]
    }
    rows.append(row_data)

print(f"有效数据行数: {len(rows)}")
print("=" * 80)

# 按模块统计
from collections import Counter
module_counter = Counter(r["所属模块"] for r in rows)
print("\n按模块统计:")
for module, count in module_counter.most_common():
    print(f"  {module}: {count}张表")

# 按是否清理统计
clean_counter = Counter(r["是否数据清理"] for r in rows)
print("\n按是否数据清理统计:")
for k, v in clean_counter.most_common():
    print(f"  {k}: {v}张表")

# 按归档策略统计
archive_counter = Counter(r["归档策略"] for r in rows)
print("\n按归档策略统计:")
for k, v in archive_counter.most_common():
    print(f"  {repr(k)}: {v}张表")

print("\n" + "=" * 80)
print("完整清单:")
print("=" * 80)
for i, r in enumerate(rows, 1):
    print(f"\n--- 第{i}条 ---")
    print(f"  序号: {r['序号']}")
    print(f"  所属数据库: {r['所属数据库']}")
    print(f"  所属模块: {r['所属模块']}")
    print(f"  责任人: {r['责任人']}")
    print(f"  表名(英文): {r['表名_英文']}")
    print(f"  表名(中文): {r['表名_中文']}")
    print(f"  是否数据清理: {r['是否数据清理']}")
    clean_strategy = r['清理策略'] if r['清理策略'] else '(空)'
    print(f"  清理策略: {clean_strategy}")
    archive_strategy = r['归档策略'] if r['归档策略'] else '(空)'
    print(f"  归档策略: {archive_strategy}")
    remark = r['备注说明'] if r['备注说明'] else '(空)'
    print(f"  备注说明: {remark}")

# 保存为JSON
output_path = r"e:\WorkSpace\HelloWorldAgentSkills\aiguibin-common-excel\cleanup_archive_list.json"
with open(output_path, 'w', encoding='utf-8') as f:
    json.dump(rows, f, ensure_ascii=False, indent=2)
print(f"\n\n完整数据已保存到: {output_path}")
