# -*- coding: utf-8 -*-
"""
重新生成维度分析文件,只保留82张大表/快速增长表
从144张中排除62张非大表
"""
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from filter_by_volume import judge_data_volume
from append_by_business_logic import generate_cleanup_strategy, generate_archive_strategy, generate_dimension_note

SRC = r'e:\WorkSpace\HelloWorldAgentSkills\aiguibin-common-excel\DD-湖北农信基础配置检核\湖北农信清理归档表清单_维度分析.xlsx'
ANALYSIS_FILE = r'e:\WorkSpace\HelloWorldAgentSkills\aiguibin-common-excel\archive_analysis_by_logic.xlsx'


def main():
    # 读取业务逻辑分析结果
    analysis_df = pd.read_excel(ANALYSIS_FILE)
    archive_df = analysis_df[analysis_df['是否需要归档'] == '是'].copy()

    print(f'需归档表总数: {len(archive_df)}')

    # 第一步: 按配额筛选144张(与之前一致)
    class_counts = archive_df.groupby('业务分类').size().to_dict()
    class_priority_weight = {
        '日志审计': 3, '营销数据': 3, '消息通知': 3,
        '风控数据': 2, '征信数据': 2, '引擎日志': 2, '批量日志': 2,
        '流水记录': 2, '历史变更': 2, '业务明细': 2,
        '贷后管理': 2, '线上贷款': 2, '押品数据': 2,
        '授信审批': 2, '用信审批': 2, '合同数据': 2,
        '放还款交易': 2, '额度数据': 2, '评级数据': 2,
        '流程数据': 2, '流程任务': 2, '档案管理': 2, '风险分类': 2,
    }

    MAX_NEW = 144
    quota = {}
    total_weight = 0
    for cls, count in class_counts.items():
        weight = class_priority_weight.get(cls, 1)
        total_weight += weight * count

    for cls, count in class_counts.items():
        weight = class_priority_weight.get(cls, 1)
        q = max(1, int(MAX_NEW * weight * count / total_weight))
        quota[cls] = min(q, count)

    total_quota = sum(quota.values())
    if total_quota > MAX_NEW:
        sorted_cls = sorted(quota.items(), key=lambda x: -x[1])
        idx = 0
        while total_quota > MAX_NEW:
            cls, q = sorted_cls[idx % len(sorted_cls)]
            if quota[cls] > 1:
                quota[cls] -= 1
                total_quota -= 1
            idx += 1
    elif total_quota < MAX_NEW:
        sorted_cls = sorted(class_counts.items(), key=lambda x: -x[1])
        idx = 0
        while total_quota < MAX_NEW:
            cls = sorted_cls[idx % len(sorted_cls)][0]
            if quota.get(cls, 0) < class_counts[cls]:
                quota[cls] = quota.get(cls, 0) + 1
                total_quota += 1
            idx += 1

    # 按配额筛选144张
    selected_144 = []
    for biz_class, q in quota.items():
        class_df = archive_df[archive_df['业务分类'] == biz_class].head(q)
        for _, row in class_df.iterrows():
            selected_144.append(row.to_dict())

    selected_144_df = pd.DataFrame(selected_144)
    print(f'第一步筛选: {len(selected_144_df)}张')

    # 第二步: 基于行业经验排除非大表
    selected_final = []
    excluded = []
    for _, row in selected_144_df.iterrows():
        table_en = str(row['表英文名'])
        table_cn = str(row['表中文名'])
        module = str(row['模块'])

        keep, volume, growth, reason = judge_data_volume(table_en, table_cn, module)

        if keep:
            selected_final.append({
                'row': row,
                'volume': volume,
                'growth': growth,
            })
        else:
            excluded.append({
                'table_en': table_en,
                'table_cn': table_cn,
                'reason': reason,
            })

    print(f'第二步筛选(排除非大表): 保留{len(selected_final)}张, 排除{len(excluded)}张')

    # 加载已有工作簿(原始86张表)
    wb = load_workbook(SRC)
    ws = wb['目录']
    last_row = ws.max_row

    # 获取已有表名
    existing_tables = set()
    for row_idx in range(2, last_row + 1):
        table_name = ws.cell(row=row_idx, column=5).value
        if table_name:
            existing_tables.add(str(table_name).strip().upper())

    priority_colors = {
        'P0': 'FFC7CE', 'P1': 'FFEB9C', 'P2': 'C6EFCE', 'P3': 'BDD7EE',
    }

    added = 0
    current_row = last_row + 1

    for item in selected_final:
        row = item['row']
        table_en = str(row['表英文名']).strip()

        if table_en.upper() in existing_tables:
            continue

        table_cn = str(row['表中文名'])
        module = str(row['模块'])
        db = str(row['数据库'])
        biz_class = str(row['业务分类'])
        volume = item['volume']
        growth = item['growth']

        cleanup_strategy = generate_cleanup_strategy(biz_class, table_cn)
        archive_strategy = generate_archive_strategy(biz_class)
        dimension_note = generate_dimension_note(biz_class, module, table_cn)

        # 在维度说明中追加数据量级别信息
        dimension_note += f'；数据量{volume}；增长趋势{growth}'

        # 写入数据
        ws.cell(row=current_row, column=1, value=added + 1)
        ws.cell(row=current_row, column=2, value=db)
        ws.cell(row=current_row, column=3, value=module)
        ws.cell(row=current_row, column=4, value='系统分析')
        ws.cell(row=current_row, column=5, value=table_en)
        ws.cell(row=current_row, column=6, value=table_cn)
        ws.cell(row=current_row, column=7, value='是')
        ws.cell(row=current_row, column=8, value=cleanup_strategy)
        ws.cell(row=current_row, column=9, value=archive_strategy)

        cell_j = ws.cell(row=current_row, column=10, value=dimension_note)
        cell_j.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

        for priority, color in priority_colors.items():
            if f'归档优先级{priority}' in dimension_note:
                cell_j.fill = PatternFill('solid', start_color=color)
                break

        ws.cell(row=current_row, column=8).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws.cell(row=current_row, column=9).alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

        added += 1
        current_row += 1

    ws.column_dimensions['H'].width = 50
    ws.column_dimensions['I'].width = 30
    ws.column_dimensions['J'].width = 70

    OUTPUT = SRC.replace('.xlsx', '_大表筛选版.xlsx')
    wb.save(OUTPUT)

    print(f'\n追加完成: 新增 {added} 张表')
    print(f'总表数: {last_row - 1 + added}张')
    print(f'输出文件: {OUTPUT}')

    # 统计
    print('\n保留表按业务分类统计:')
    biz_counts = {}
    for item in selected_final:
        bc = item['row']['业务分类']
        biz_counts[bc] = biz_counts.get(bc, 0) + 1
    for bc, c in sorted(biz_counts.items(), key=lambda x: -x[1]):
        print(f'  {bc}: {c}张')

    print(f'\n排除表按排除理由统计:')
    reason_counts = {}
    for item in excluded:
        r = item['reason']
        reason_counts[r] = reason_counts.get(r, 0) + 1
    for r, c in sorted(reason_counts.items(), key=lambda x: -x[1]):
        print(f'  {r}: {c}张')


if __name__ == '__main__':
    main()
