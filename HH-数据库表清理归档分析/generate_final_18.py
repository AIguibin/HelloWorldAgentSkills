# -*- coding: utf-8 -*-
"""
只保留用户指定的18张表,与原始86张表合并生成最终文件
"""
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

SRC = r'e:\WorkSpace\HelloWorldAgentSkills\aiguibin-common-excel\DD-湖北农信基础配置检核\湖北农信清理归档表清单_维度分析.xlsx'
BIZ_LOGIC_FILE = r'e:\WorkSpace\HelloWorldAgentSkills\aiguibin-common-excel\DD-湖北农信基础配置检核\湖北农信清理归档表清单_维度分析_业务逻辑版.xlsx'

# 用户指定的18张表(中文名)
SPECIFIED_TABLES_CN = [
    '个人客户变更明细表',
    '客户信息变更明细表',
    '对公客户_企业电费缴纳明细表',
    '对公客户_企业燃气费缴纳明细表',
    '对公客户_企业水费缴纳明细表',
    '产品管理_产品要素_明细发布历史表',
    '产品管理_产品管理_信息发布历史表',
    '变量修改历史记录表',
    '对公客户变更记录表',
    '个人客户变更记录表',
    '标签历史数据更新记录表',
    '客户管理_集群客户_变更历史表',
    'rule_log_details',
    'rule_log_report_count',
    'rule_logs_count',
    'rule_logs_report',
    'rule_logsbackup',
    '对公客户变更明细记录表',
]


def main():
    # 读取业务逻辑版文件(含144张新增表)
    biz_df = pd.read_excel(BIZ_LOGIC_FILE, sheet_name='目录')
    new_df = biz_df.iloc[86:].copy()

    print(f'业务逻辑版新增表数: {len(new_df)}')

    # 筛选用户指定的18张表
    # 按中文名匹配(第6列,索引5)
    selected = new_df[new_df.iloc[:, 5].astype(str).isin(SPECIFIED_TABLES_CN)]
    # 按英文名匹配(第5列,索引4)
    selected_en = new_df[new_df.iloc[:, 4].astype(str).isin(SPECIFIED_TABLES_CN)]
    # 合并去重
    selected = pd.concat([selected, selected_en]).drop_duplicates(subset=new_df.columns[4])

    print(f'匹配到指定表数: {len(selected)}')
    print('\n匹配的表清单:')
    for _, row in selected.iterrows():
        print(f'  [{row.iloc[2]}] {row.iloc[4]} | {row.iloc[5]}')

    # 检查是否有未匹配的表
    matched_cn = set(selected.iloc[:, 5].astype(str))
    matched_en = set(selected.iloc[:, 4].astype(str))
    matched = matched_cn | matched_en
    not_found = [t for t in SPECIFIED_TABLES_CN if t not in matched]
    if not_found:
        print(f'\n未匹配的表: {not_found}')

    # 加载原始86张表的工作簿
    wb = load_workbook(SRC)
    ws = wb['目录']
    last_row = ws.max_row

    # 获取已有表名
    existing_tables = set()
    for row_idx in range(2, last_row + 1):
        table_name = ws.cell(row=row_idx, column=5).value
        if table_name:
            existing_tables.add(str(table_name).strip().upper())

    priority_colors = {
        'P0': 'FFC7CE', 'P1': 'FFEB9C', 'P2': 'C6EFCE', 'P3': 'BDD7EE',
    }

    added = 0
    current_row = last_row + 1

    for _, row in selected.iterrows():
        table_en = str(row.iloc[4]).strip()
        if table_en.upper() in existing_tables:
            continue

        # 复制所有列数据
        for col_idx in range(1, 11):  # A-J列
            value = row.iloc[col_idx - 1]
            if pd.isna(value):
                value = None
            cell = ws.cell(row=current_row, column=col_idx, value=value)

            # J列(维度说明)设置格式
            if col_idx == 10:
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
                dimension_note = str(value) if value else ''
                for priority, color in priority_colors.items():
                    if f'归档优先级{priority}' in dimension_note:
                        cell.fill = PatternFill('solid', start_color=color)
                        break
            elif col_idx in (8, 9):  # H列、I列
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

        added += 1
        current_row += 1

    # 调整列宽
    ws.column_dimensions['H'].width = 50
    ws.column_dimensions['I'].width = 30
    ws.column_dimensions['J'].width = 70

    # 保存
    OUTPUT = SRC.replace('.xlsx', '_最终版.xlsx')
    wb.save(OUTPUT)

    print(f'\n追加完成: 新增 {added} 张表')
    print(f'总表数: {last_row - 1 + added}张')
    print(f'输出文件: {OUTPUT}')


if __name__ == '__main__':
    main()
